"""Admin-only dashboard endpoints."""

import logging
import re
import zipfile
from io import BytesIO
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Literal, Optional

from beanie import PydanticObjectId
from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from PIL import Image as PILImage
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# Executive report table theme (reference: gray label column, outside border only).
_PPTX_TABLE_FONT = "Trebuchet MS"
_PPTX_TABLE_FONT_PT = 8
_PPTX_LABEL_FILL = RGBColor(166, 166, 166)  # medium gray
_PPTX_LABEL_TEXT = RGBColor(255, 255, 255)
_PPTX_VALUE_FILL = RGBColor(255, 255, 255)
_PPTX_VALUE_TEXT = RGBColor(0, 0, 0)
_PPTX_TABLE_BORDER_HEX = "454545"
from pydantic import BaseModel, Field

from models import Defect, User, UserLog
from prompt import build_executive_defect_report_prompt
from services.generate_client import (
    classify_construction_site_image,
    generate_executive_defect_report,
)
from services.report_defect_extract import estimate_wrapped_line_count, invalid_image_fields
from utils.deps import require_admin
from utils.security import hash_password

logger = logging.getLogger(__name__)

REPO_ROOT = Path(__file__).resolve().parent.parent.parent
UPLOADS_DIR = REPO_ROOT / "uploads"

router = APIRouter(prefix="/api/admin", tags=["admin"])


class AdminResetPasswordBody(BaseModel):
    new_password: str = Field(..., min_length=8, max_length=128)


class AdminRoleBody(BaseModel):
    role: Literal["user", "admin"]


class ReportAnalyzeItemBody(BaseModel):
    defect_id: str


class ReportXlsxEntry(BaseModel):
    defect_id: str
    observation: str = ""
    recommendation: str = ""


class ReportGenerateXlsxBody(BaseModel):
    entries: list[ReportXlsxEntry]
    base_url: str = ""


class BulkDeleteUploadsBody(BaseModel):
    """Bulk-remove defect uploads (files + DB rows)."""

    mode: Literal["all", "user", "ids"]
    user_id: Optional[str] = None
    defect_ids: list[str] = Field(default_factory=list)


async def _get_user_or_404(user_id: str) -> User:
    try:
        user = await User.get(PydanticObjectId(user_id))
    except Exception:
        raise HTTPException(status_code=400, detail="Invalid user ID") from None
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    return user


async def _log_admin_action(
    admin: User,
    action: str,
    *,
    target: User,
) -> None:
    await UserLog(
        user_id=str(admin.id),
        action=action,
        target_user_id=str(target.id),
        target_email=target.email,
    ).insert()


async def _purge_user_defect_files(user_id: str) -> int:
    """Delete defect rows and image files for a user. Returns number of defects removed."""
    defects = await Defect.find(Defect.user_id == user_id).to_list()
    removed = 0
    for defect in defects:
        await _delete_defect_upload_best_effort(defect)
        removed += 1
    return removed


async def _delete_defect_upload_best_effort(defect: Defect) -> None:
    """Remove DB row; delete file when possible without blocking row removal."""
    file_path = REPO_ROOT / defect.image_path
    if file_path.exists():
        try:
            file_path.unlink()
        except OSError as e:
            logger.error("admin: could not delete file %s: %s", file_path, e)
    await defect.delete()


async def _delete_defect_upload_strict(defect: Defect) -> None:
    """Remove one defect's image file (if present) and DB row; mirrors DELETE /uploads/{id}."""
    file_path = REPO_ROOT / defect.image_path
    if file_path.exists():
        try:
            file_path.unlink()
        except OSError as e:
            logger.error(
                "admin: failed to delete file %s for upload %s: %s",
                file_path, defect.id, e,
            )
            raise HTTPException(
                status_code=500,
                detail="Failed to delete file from storage. Database record preserved.",
            ) from e
    await defect.delete()


async def _delete_user_profile_photo(user: User) -> None:
    if not user.profile_photo or not user.profile_photo.startswith("uploads/"):
        return
    path = REPO_ROOT / user.profile_photo
    try:
        path.unlink(missing_ok=True)
    except OSError as e:
        logger.warning("admin: profile photo unlink failed %s: %s", path, e)


def _is_user_enabled(u: User) -> bool:
    return not bool(getattr(u, "is_disabled", False))


async def _count_enabled_admins() -> int:
    admins = await User.find(User.role == "admin").to_list()
    return sum(1 for u in admins if _is_user_enabled(u))


async def _count_total_admins() -> int:
    return await User.find(User.role == "admin").count()


def _as_utc_aware(dt: datetime | None) -> datetime | None:
    """Mongo often returns naive datetimes that are UTC; JS needs an offset or Z to parse correctly."""
    if dt is None:
        return None
    if dt.tzinfo is None:
        return dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(timezone.utc)


def _iso_utc(dt: datetime | None) -> str | None:
    d = _as_utc_aware(dt)
    if d is None:
        return None
    # e.g. 2026-04-28T12:00:00.123+00:00 — works with new Date() in browsers
    return d.isoformat(timespec="milliseconds")


def _defect_to_dict(d: Defect, *, email: str | None = None, name: str | None = None) -> dict:
    out = {
        "id": str(d.id),
        "user_id": d.user_id,
        "image_path": d.image_path,
        "project": getattr(d, "project", "") or "",
        "tower": d.tower,
        "floor": d.floor,
        "flat": d.flat,
        "room": d.room,
        "category": d.category,
        "description": d.description,
        "created_at": _iso_utc(_as_utc_aware(d.created_at)),
    }
    if email is not None:
        out["email"] = email
    if name is not None:
        out["name"] = name
    return out


def _defect_location_string(d: Defect) -> str:
    project = str(getattr(d, "project", "") or "").strip()
    base = f"Tower {d.tower}, Floor {d.floor}, Flat {d.flat}, Room {d.room}"
    if project:
        return f"{project}, {base}"
    return base


def _pptx_bullet_display(value: str) -> str:
    """Use stored observation/recommendation text as-is (already validated bullets)."""
    text = str(value or "").strip()
    if not text:
        return "To be confirmed"
    if text.lower().startswith("invalid image detected"):
        return text
    return text


_PPTX_MONTH_ABBR = (
    "Jan",
    "Feb",
    "Mar",
    "Apr",
    "May",
    "Jun",
    "Jul",
    "Aug",
    "Sep",
    "Oct",
    "Nov",
    "Dec",
)


def _pptx_format_reporting_date(value: str) -> str:
    """Display reporting date as e.g. 21-May-2026 (day-Mon-year)."""
    raw = str(value or "").strip()
    if not raw:
        return "To be confirmed"
    if raw.lower() == "to be confirmed":
        return raw

    parsed: date | None = None
    if len(raw) >= 10 and raw[4:5] == "-" and raw[7:8] == "-":
        try:
            parsed = date.fromisoformat(raw[:10])
        except ValueError:
            parsed = None
    if parsed is None and "T" in raw:
        try:
            parsed = datetime.fromisoformat(raw.replace("Z", "+00:00")).date()
        except ValueError:
            parsed = None
    if parsed is None:
        return raw
    return f"{parsed.day}-{_PPTX_MONTH_ABBR[parsed.month - 1]}-{parsed.year}"


def _pptx_value_wrap_chars(card_w_in: float, label_col_w_in: float) -> int:
    value_w = max(1.4, card_w_in - label_col_w_in)
    return max(28, int(value_w * 15))


def _pptx_row_heights_for_card(
    row: dict,
    *,
    value_chars: int,
    meta_h: float = 0.27,
    line_h: float = 0.135,
) -> tuple[float, float, float, float, float, str, str]:
    """Return meta/obs/rec row heights (inches) and display strings."""
    loc = _report_card_location(row) or "To be confirmed"
    category = str(row.get("category") or "").strip() or "To be confirmed"
    reporting_date = _pptx_format_reporting_date(str(row.get("reporting_date") or ""))
    obs_field = _pptx_bullet_display(row.get("observation") or "")
    rec_field = _pptx_bullet_display(row.get("recommendation") or "")

    loc_h = max(meta_h, 0.2 + estimate_wrapped_line_count(loc, chars_per_line=value_chars) * line_h)
    cat_h = max(meta_h, 0.2 + estimate_wrapped_line_count(category, chars_per_line=value_chars) * line_h)
    date_h = max(meta_h, 0.2 + estimate_wrapped_line_count(reporting_date, chars_per_line=value_chars) * line_h)
    obs_h = max(0.32, 0.2 + estimate_wrapped_line_count(obs_field, chars_per_line=value_chars) * line_h)
    rec_h = max(0.32, 0.2 + estimate_wrapped_line_count(rec_field, chars_per_line=value_chars) * line_h)

    return loc_h, cat_h, date_h, obs_h, rec_h, obs_field, rec_field


def _pptx_slide_table_budget_in() -> float:
    """Max table stack height (inches) before footer on 7.5\" slide."""
    return 2.95


# Excel report image column — embedded previews (pixels at 96 DPI).
_EXCEL_IMAGE_MAX_W_PX = 400
_EXCEL_IMAGE_MAX_H_PX = 300
_EXCEL_IMAGE_CELL_PAD_PX = 10
_EXCEL_IMAGE_COL_MIN_WIDTH = 58.0
_EXCEL_IMAGE_JPEG_QUALITY = 92
_EXCEL_ROW_HEIGHT_FACTOR = 0.75  # user-facing row height vs raw content box


def _excel_pixels_to_points(px: float) -> float:
    return float(px) * 72.0 / 96.0


def _excel_col_chars_to_pixels(width_chars: float) -> int:
    """Approximate pixel width for default Calibri 11 column width."""
    return max(1, int(float(width_chars) * 7 + 5))


def _excel_row_points_to_pixels(height_pt: float) -> int:
    return max(1, int(float(height_pt) * 96 / 72))


def _excel_column_width_for_pixels(px: int) -> float:
    """Approximate Excel column width (character units) for embedded image pixels."""
    return round(max(_EXCEL_IMAGE_COL_MIN_WIDTH, (int(px) + 14) / 7.0), 2)


def _excel_row_height_for_content(*, image_h_px: float, text_height_pt: float) -> float:
    """Row height: text may use 0.75 factor; image rows must fit the full thumbnail height."""
    pad_pt = _excel_pixels_to_points(_EXCEL_IMAGE_CELL_PAD_PX * 2)
    text_row_pt = max(24.0, text_height_pt) * _EXCEL_ROW_HEIGHT_FACTOR
    if image_h_px <= 0:
        return max(20.0, text_row_pt)
    image_block_pt = _excel_pixels_to_points(image_h_px) + pad_pt
    # +2pt avoids sub-pixel rounding letting drawings bleed into the next row.
    return max(20.0, text_row_pt, image_block_pt) + 2.0


def _excel_fit_image_in_cell_bounds(
    embed_w: int,
    embed_h: int,
    col_w_px: int,
    row_h_px: int,
) -> tuple[int, int, int, int]:
    """Scale image to fit inside the cell box; return display size and centering offsets."""
    pad = _EXCEL_IMAGE_CELL_PAD_PX
    inner_w = max(1, col_w_px - (2 * pad))
    inner_h = max(1, row_h_px - (2 * pad))
    scale = min(1.0, inner_w / max(embed_w, 1), inner_h / max(embed_h, 1))
    display_w = max(1, int(embed_w * scale))
    display_h = max(1, int(embed_h * scale))
    col_off, row_off = _excel_center_offsets(
        col_w_px,
        row_h_px,
        display_w,
        display_h,
        min_pad=pad,
    )
    return display_w, display_h, col_off, row_off


def _excel_prepare_image_embed(path: Path) -> tuple[int, int, BytesIO] | None:
    """Resize source photo for Excel (upscale small images, high JPEG quality)."""
    try:
        with PILImage.open(path) as img:
            img = img.convert("RGB")
            w, h = img.size
            scale = min(
                _EXCEL_IMAGE_MAX_W_PX / max(w, 1),
                _EXCEL_IMAGE_MAX_H_PX / max(h, 1),
            )
            new_w = max(1, int(w * scale))
            new_h = max(1, int(h * scale))
            if (new_w, new_h) != (w, h):
                img = img.resize((new_w, new_h), PILImage.Resampling.LANCZOS)
            out = BytesIO()
            img.save(out, format="JPEG", optimize=True, quality=_EXCEL_IMAGE_JPEG_QUALITY)
            out.seek(0)
            return new_w, new_h, out
    except Exception:
        logger.exception("admin report: failed to prepare excel image %s", path)
        return None


def _excel_center_offsets(
    cell_w_px: int,
    cell_h_px: int,
    img_w: int,
    img_h: int,
    *,
    min_pad: int = _EXCEL_IMAGE_CELL_PAD_PX,
) -> tuple[int, int]:
    inner_w = max(0, cell_w_px - (2 * min_pad))
    inner_h = max(0, cell_h_px - (2 * min_pad))
    col_off = min_pad + max(0, (inner_w - img_w) // 2)
    row_off = min_pad + max(0, (inner_h - img_h) // 2)
    return col_off, row_off


def _excel_prepare_row_images(rows: list[dict]) -> tuple[list[tuple[int, int, BytesIO | None]], int]:
    """Pre-render embed-sized JPEGs and display dimensions for column H."""
    plans: list[tuple[int, int, BytesIO | None]] = []
    max_embed_w_px = 0
    for row in rows:
        abs_img = _safe_image_abs_path(str(row.get("image_path") or ""))
        if not abs_img:
            plans.append((0, 0, None))
            continue
        prepared = _excel_prepare_image_embed(abs_img)
        if prepared is None:
            plans.append((0, 0, None))
            continue
        embed_w, embed_h, thumb = prepared
        plans.append((embed_w, embed_h, thumb))
        max_embed_w_px = max(max_embed_w_px, embed_w)
    return plans, max_embed_w_px


def _excel_finalize_image_column(ws, *, col: int = 8, first_row: int = 2) -> None:
    """Ensure Image column cells have no value or hyperlink metadata."""
    col_letter = get_column_letter(col)
    for row_idx in range(first_row, (ws.max_row or first_row) + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.value = None
        cell.hyperlink = None
    hyperlinks = list(getattr(ws, "_hyperlinks", []) or [])
    if hyperlinks:
        ws._hyperlinks = [
            h
            for h in hyperlinks
            if not str(getattr(h, "ref", "")).upper().startswith(col_letter)
        ]


def _scrub_excel_column_h_hyperlinks_zip(xlsx_buf: BytesIO) -> BytesIO:
    """
    Remove column-H hyperlinks and path-like cell values from worksheet XML.
    Excel shows hyperlink targets as visible text even when the cell looks empty.
    """
    hyperlink_re = re.compile(rb'<hyperlink\b[^>]*\bref="H\d+"[^>]*/>', re.IGNORECASE)
    h_cell_value_re = re.compile(
        rb'(<c r="H\d+"[^>]*>(?:(?!</c>).)*?<v>)([^<]*)(</v>)',
        re.IGNORECASE | re.DOTALL,
    )

    def _strip_path_values(match: re.Match[bytes]) -> bytes:
        value = match.group(2)
        lower = value.lower()
        if (
            b"/" in value
            or b"http" in lower
            or b"uploads" in lower
            or b".jpg" in lower
            or b".jpeg" in lower
            or b".png" in lower
            or b".webp" in lower
        ):
            return match.group(1) + match.group(3)
        return match.group(0)

    xlsx_buf.seek(0)
    out_buf = BytesIO()
    with zipfile.ZipFile(xlsx_buf, "r") as zin, zipfile.ZipFile(out_buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/") and item.filename.endswith(".xml"):
                data = hyperlink_re.sub(b"", data)
                data = h_cell_value_re.sub(_strip_path_values, data)
            zout.writestr(item, data)
    out_buf.seek(0)
    return out_buf


def _excel_embed_image_in_cell(
    ws,
    xl_img: XLImage,
    *,
    row: int,
    col: int = 8,
    col_off_px: int = _EXCEL_IMAGE_CELL_PAD_PX,
    row_off_px: int = _EXCEL_IMAGE_CELL_PAD_PX,
) -> None:
    """Anchor image inside the cell; offsets center it within the row/column box."""
    xl_img.anchor = OneCellAnchor(
        _from=AnchorMarker(
            col=col - 1,
            row=row - 1,
            colOff=pixels_to_EMU(col_off_px),
            rowOff=pixels_to_EMU(row_off_px),
        ),
        ext=XDRPositiveSize2D(
            cx=pixels_to_EMU(int(xl_img.width)),
            cy=pixels_to_EMU(int(xl_img.height)),
        ),
    )
    ws.add_image(xl_img)


def _safe_image_abs_path(image_path: str) -> Path | None:
    rel = str(image_path or "").replace("\\", "/").lstrip("/")
    if not rel:
        return None
    abs_path = (REPO_ROOT / rel).resolve()
    try:
        abs_path.relative_to(REPO_ROOT.resolve())
    except ValueError:
        return None
    if not abs_path.exists() or not abs_path.is_file():
        return None
    return abs_path


def _thumbnail_image_for_pptx(
    path: Path,
    *,
    max_w: int = 1400,
    max_h: int = 1000,
    jpeg_quality: int = 84,
) -> tuple[BytesIO, int, int] | None:
    try:
        with PILImage.open(path) as img:
            img = img.convert("RGB")
            img.thumbnail((max_w, max_h), PILImage.Resampling.LANCZOS)
            out = BytesIO()
            img.save(out, format="JPEG", optimize=True, quality=jpeg_quality)
            out.seek(0)
            return out, img.width, img.height
    except Exception:
        logger.exception("admin report: failed to create pptx image for %s", path)
        return None


def _report_card_location(row: dict) -> str:
    tower = str(row.get("tower") or "").strip()
    floor = str(row.get("floor") or "").strip()
    flat = str(row.get("flat") or "").strip()
    room = str(row.get("room") or "").strip()

    parts: list[str] = []
    if tower:
        parts.append(f"T-{tower}" if not tower.upper().startswith("T") else tower)
    if floor:
        parts.append(f"Floor {floor}")
    if flat:
        parts.append(f"Flat {flat}")
    if room:
        parts.append(room)
    return ", ".join(parts) or "To be confirmed"


def _apply_pptx_cell_surface(cell, *, label: bool = False) -> None:
    fill = cell.fill
    fill.solid()
    fill.fore_color.rgb = _PPTX_LABEL_FILL if label else _PPTX_VALUE_FILL


def _clear_pptx_cell_borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        for existing in tc_pr.findall(qn(f"a:{edge}")):
            tc_pr.remove(existing)


def _set_pptx_cell_border_edges(
    cell,
    *,
    left: bool = False,
    right: bool = False,
    top: bool = False,
    bottom: bool = False,
    color_hex: str = _PPTX_TABLE_BORDER_HEX,
) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for edge, enabled in (
        ("lnL", left),
        ("lnR", right),
        ("lnT", top),
        ("lnB", bottom),
    ):
        if not enabled:
            continue
        for existing in tc_pr.findall(qn(f"a:{edge}")):
            tc_pr.remove(existing)
        tc_pr.append(
            parse_xml(
                f'<a:{edge} xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
                f'w="12700" cap="flat" cmpd="sng" algn="ctr">'
                f'<a:solidFill><a:srgbClr val="{color_hex}"/></a:solidFill>'
                f'<a:prstDash val="solid"/></a:{edge}>'
            )
        )


def _apply_pptx_table_outside_borders(table, *, rows: int, cols: int = 2) -> None:
    """Single outside frame — no internal grid lines between rows or columns."""
    for r in range(rows):
        for c in range(cols):
            _clear_pptx_cell_borders(table.cell(r, c))
    for r in range(rows):
        for c in range(cols):
            _set_pptx_cell_border_edges(
                table.cell(r, c),
                left=c == 0,
                right=c == cols - 1,
                top=r == 0,
                bottom=r == rows - 1,
            )


def _pptx_set_run_font(run, *, label: bool) -> None:
    run.font.name = _PPTX_TABLE_FONT
    run.font.size = Pt(_PPTX_TABLE_FONT_PT)
    run.font.bold = False
    run.font.color.rgb = _PPTX_LABEL_TEXT if label else _PPTX_VALUE_TEXT


def _set_pptx_cell(
    cell,
    text: str,
    *,
    label: bool = False,
    multiline: bool = False,
) -> None:
    _apply_pptx_cell_surface(cell, label=label)
    cell.vertical_anchor = MSO_ANCHOR.TOP
    cell.margin_left = Inches(0.08 if label else 0.07)
    cell.margin_right = Inches(0.06)
    cell.margin_top = Inches(0.05)
    cell.margin_bottom = Inches(0.05)

    tf = cell.text_frame
    tf.word_wrap = True
    tf.clear()
    content = str(text or "To be confirmed")
    lines = content.split("\n") if multiline else [content]

    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(0)
        p.space_after = Pt(2 if multiline and i < len(lines) - 1 else 0)
        p.line_spacing = 1.0
        run = p.add_run()
        run.text = line
        _pptx_set_run_font(run, label=label)


def _add_pptx_observation_badge(slide, left, top, number: int) -> None:
    """Small bordered badge at top-left of the image frame."""
    pad = Inches(0.1)
    badge_w = Inches(0.4)
    badge_h = Inches(0.24)
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        int(left + pad),
        int(top + pad),
        int(badge_w),
        int(badge_h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(148, 163, 184)
    shape.line.width = Pt(0.85)

    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = str(number)
    run.font.name = "Aptos"
    run.font.size = Pt(9)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 41, 59)


def _add_fitted_pptx_image(slide, path: Path, left, top, width, height) -> None:
    image_data = _thumbnail_image_for_pptx(path)
    if image_data is None:
        placeholder = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        placeholder.fill.solid()
        placeholder.fill.fore_color.rgb = RGBColor(248, 250, 252)
        placeholder.line.color.rgb = RGBColor(203, 213, 225)
        placeholder.text = "Image unavailable"
        return

    image_stream, px_w, px_h = image_data
    scale = min(width / max(px_w, 1), height / max(px_h, 1))
    actual_w = int(px_w * scale)
    actual_h = int(px_h * scale)
    img_left = int(left + (width - actual_w) / 2)
    img_top = int(top + (height - actual_h) / 2)
    slide.shapes.add_picture(image_stream, img_left, img_top, width=actual_w, height=actual_h)


def _build_report_presentation(*, rows: list[dict]) -> BytesIO:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    title_color = RGBColor(190, 54, 73)
    subtitle_color = RGBColor(156, 43, 59)
    image_border_color = RGBColor(69, 69, 69)

    card_w = Inches(3.58)
    card_w_in = 3.58
    gap = Inches(0.36)
    start_x = Inches(0.7)
    image_top_in = 1.05
    image_h_in = 2.92
    table_gap_in = 0.08
    img_pad = Inches(0.05)
    label_col_w = Inches(1.34)
    label_col_w_in = 1.34
    value_chars = _pptx_value_wrap_chars(card_w_in, label_col_w_in)
    table_budget_in = _pptx_slide_table_budget_in()

    for slide_start in range(0, len(rows), 3):
        chunk = rows[slide_start : slide_start + 3]
        slide = prs.slides.add_slide(blank_layout)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)

        card_plans: list[dict] = []
        max_table_h_in = 0.0
        for row in chunk:
            loc_h, cat_h, date_h, obs_h, rec_h, obs_field, rec_field = _pptx_row_heights_for_card(
                row, value_chars=value_chars
            )
            total_h = loc_h + cat_h + date_h + obs_h + rec_h
            max_table_h_in = max(max_table_h_in, total_h)
            card_plans.append(
                {
                    "row": row,
                    "loc_h": loc_h,
                    "cat_h": cat_h,
                    "date_h": date_h,
                    "obs_h": obs_h,
                    "rec_h": rec_h,
                    "obs_field": obs_field,
                    "rec_field": rec_field,
                    "total_h": total_h,
                }
            )

        slide_image_h_in = image_h_in
        if max_table_h_in > table_budget_in:
            slide_image_h_in = max(2.1, image_h_in - (max_table_h_in - table_budget_in) - 0.06)

        image_top = Inches(image_top_in)
        image_h = Inches(slide_image_h_in)
        table_top = Inches(image_top_in + slide_image_h_in + table_gap_in)

        title_box = slide.shapes.add_textbox(Inches(0.55), Inches(0.18), Inches(8.7), Inches(0.36))
        title_tf = title_box.text_frame
        title_tf.clear()
        title_p = title_tf.paragraphs[0]
        title_run = title_p.add_run()
        title_run.text = "Quality walkthrough - Observation"
        title_run.font.name = "Aptos"
        title_run.font.size = Pt(21)
        title_run.font.italic = True
        title_run.font.bold = True
        title_run.font.color.rgb = title_color

        categories = {str(r.get("category") or "Observation").strip() for r in chunk if r.get("category")}
        subtitle = next(iter(categories)) if len(categories) == 1 else "Defect observations"
        sub_box = slide.shapes.add_textbox(Inches(0.58), Inches(0.56), Inches(8.2), Inches(0.26))
        sub_tf = sub_box.text_frame
        sub_tf.clear()
        sub_run = sub_tf.paragraphs[0].add_run()
        sub_run.text = subtitle
        sub_run.font.name = "Aptos"
        sub_run.font.size = Pt(11)
        sub_run.font.italic = True
        sub_run.font.bold = True
        sub_run.font.color.rgb = subtitle_color

        for idx, plan in enumerate(card_plans):
            row = plan["row"]
            left = start_x + idx * (card_w + gap)

            img_border = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, image_top, card_w, image_h)
            img_border.fill.background()
            img_border.line.color.rgb = image_border_color
            img_border.line.width = Pt(0.75)

            obs_num = slide_start + idx + 1
            abs_img = _safe_image_abs_path(str(row.get("image_path") or ""))
            if abs_img:
                _add_fitted_pptx_image(
                    slide,
                    abs_img,
                    left + img_pad,
                    image_top + img_pad,
                    card_w - (img_pad * 2),
                    image_h - (img_pad * 2),
                )
            _add_pptx_observation_badge(slide, left, image_top, obs_num)

            loc = _report_card_location(row) or "To be confirmed"
            category = str(row.get("category") or "").strip() or "To be confirmed"
            reporting_date = _pptx_format_reporting_date(str(row.get("reporting_date") or ""))
            table_rows = [
                ("Location", loc),
                ("Category", category),
                ("Reporting date", reporting_date),
                ("Observation", plan["obs_field"]),
                ("Recommendation", plan["rec_field"]),
            ]
            row_heights = (
                plan["loc_h"],
                plan["cat_h"],
                plan["date_h"],
                plan["obs_h"],
                plan["rec_h"],
            )
            table_shape = slide.shapes.add_table(
                5, 2, left, table_top, card_w, Inches(plan["total_h"])
            )
            table = table_shape.table
            table.columns[0].width = label_col_w
            table.columns[1].width = card_w - label_col_w
            for row_idx, (label, value) in enumerate(table_rows):
                table.rows[row_idx].height = Inches(row_heights[row_idx])
                _set_pptx_cell(table.cell(row_idx, 0), label, label=True)
                is_bullets = row_idx >= 3
                _set_pptx_cell(
                    table.cell(row_idx, 1),
                    str(value),
                    multiline=is_bullets or "\n" in str(value),
                )
            _apply_pptx_table_outside_borders(table, rows=len(table_rows))

        footer = slide.shapes.add_textbox(Inches(0.22), Inches(7.12), Inches(3.4), Inches(0.18))
        footer_run = footer.text_frame.paragraphs[0].add_run()
        footer_run.text = "SiteSureLabs"
        footer_run.font.name = "Aptos"
        footer_run.font.size = Pt(7)
        footer_run.font.color.rgb = RGBColor(71, 85, 105)

    output = BytesIO()
    prs.save(output)
    output.seek(0)
    return output


def _build_report_workbook(
    *,
    rows: list[dict],
    base_url: str,
) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inspection Report"

    headers = [
        "S.No",
        "Category",
        "Reporting Date",
        "Tower",
        "Floor",
        "Flat",
        "Room",
        "Image",
        "Observation",
        "Recommendation",
    ]
    ws.append(headers)

    # Fixed report header color (consistent across every generated Excel).
    header_fill = PatternFill(start_color="E7C2CB", end_color="E7C2CB", fill_type="solid")
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    thin = Side(style="thin", color="E2E8F0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color="1E293B")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 24
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 10
    ws.column_dimensions["F"].width = 10
    ws.column_dimensions["G"].width = 16
    ws.column_dimensions["I"].width = 60
    ws.column_dimensions["J"].width = 60

    image_col_idx = 8
    image_plans, max_embed_w_px = _excel_prepare_row_images(rows)
    col_w_chars = _excel_column_width_for_pixels(
        (max_embed_w_px + (_EXCEL_IMAGE_CELL_PAD_PX * 2))
        if max_embed_w_px
        else (_EXCEL_IMAGE_MAX_W_PX + (_EXCEL_IMAGE_CELL_PAD_PX * 2)),
    )
    ws.column_dimensions["H"].width = col_w_chars
    col_w_px = _excel_col_chars_to_pixels(col_w_chars)

    current_row = 2
    for idx, row in enumerate(rows, start=1):
        ws.cell(row=current_row, column=1, value=idx)
        ws.cell(row=current_row, column=2, value=row.get("category") or "Others")
        ws.cell(row=current_row, column=3, value=row.get("reporting_date") or "")
        ws.cell(row=current_row, column=4, value=row.get("tower") or "")
        ws.cell(row=current_row, column=5, value=row.get("floor") or "")
        ws.cell(row=current_row, column=6, value=row.get("flat") or "")
        ws.cell(row=current_row, column=7, value=row.get("room") or "")

        image_path = str(row.get("image_path") or "")
        img_cell = ws.cell(row=current_row, column=image_col_idx, value=None)
        img_cell.hyperlink = None
        img_cell.number_format = "General"
        img_cell.font = Font(color="1E293B", underline=None)
        img_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        img_cell.fill = alt_fill if idx % 2 == 0 else PatternFill(
            start_color="FFFFFF",
            end_color="FFFFFF",
            fill_type="solid",
        )

        obs_value = str(row.get("observation") or "").strip()
        rec_value = str(row.get("recommendation") or "").strip()
        ws.cell(row=current_row, column=9, value=obs_value)
        ws.cell(row=current_row, column=10, value=rec_value)

        for col_idx in (1, 2, 3, 4, 5, 6, 7):
            c = ws.cell(row=current_row, column=col_idx)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = border
            if idx % 2 == 0:
                c.fill = alt_fill
        ws.cell(row=current_row, column=9).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.cell(row=current_row, column=10).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.cell(row=current_row, column=8).border = border
        ws.cell(row=current_row, column=9).border = border
        ws.cell(row=current_row, column=10).border = border
        if idx % 2 == 0:
            ws.cell(row=current_row, column=9).fill = alt_fill
            ws.cell(row=current_row, column=10).fill = alt_fill

        obs_text = str(ws.cell(row=current_row, column=9).value or "")
        rec_text = str(ws.cell(row=current_row, column=10).value or "")
        text_lines = max(
            estimate_wrapped_line_count(obs_text, chars_per_line=55),
            estimate_wrapped_line_count(rec_text, chars_per_line=55),
        )
        text_height_pt = 12 + (text_lines * 15)
        embed_w_px, embed_h_px, thumb_io = image_plans[idx - 1]
        row_height_pt = _excel_row_height_for_content(
            image_h_px=embed_h_px,
            text_height_pt=text_height_pt,
        )
        ws.row_dimensions[current_row].height = row_height_pt

        if thumb_io is not None and embed_w_px > 0 and embed_h_px > 0:
            try:
                thumb_io.seek(0)
                xl_img = XLImage(thumb_io)
                row_h_px = _excel_row_points_to_pixels(row_height_pt)
                display_w, display_h, col_off_px, row_off_px = _excel_fit_image_in_cell_bounds(
                    embed_w_px,
                    embed_h_px,
                    col_w_px,
                    row_h_px,
                )
                xl_img.width = display_w
                xl_img.height = display_h
                _excel_embed_image_in_cell(
                    ws,
                    xl_img,
                    row=current_row,
                    col=image_col_idx,
                    col_off_px=col_off_px,
                    row_off_px=row_off_px,
                )
            except Exception:
                logger.exception(
                    "admin report: failed to embed image %s",
                    image_path,
                )

        current_row += 1

    _excel_finalize_image_column(ws, col=image_col_idx, first_row=2)

    output = BytesIO()
    wb.save(output)
    output = _scrub_excel_column_h_hyperlinks_zip(output)
    return output


async def _collect_report_rows(entries: list[ReportXlsxEntry]) -> list[dict]:
    if not entries:
        raise HTTPException(status_code=400, detail="No report entries provided")

    try:
        ids = [PydanticObjectId(e.defect_id) for e in entries]
    except Exception:
        raise HTTPException(status_code=400, detail="Invalid defect IDs in payload") from None

    defects = await Defect.find({"_id": {"$in": ids}}).to_list()
    defect_map = {str(d.id): d for d in defects}
    rows: list[dict] = []
    for entry in entries:
        defect = defect_map.get(entry.defect_id)
        if not defect:
            continue
        rows.append(
            {
                "defect_id": str(defect.id),
                "category": defect.category or "Others",
                "reporting_date": defect.created_at.date().isoformat(),
                "tower": defect.tower,
                "floor": defect.floor,
                "flat": defect.flat,
                "room": defect.room,
                "image_path": defect.image_path,
                "observation": entry.observation,
                "recommendation": entry.recommendation,
            }
        )

    if not rows:
        raise HTTPException(status_code=404, detail="No matching uploads found for report generation")
    return rows


@router.get("/users")
async def get_all_users(_: User = Depends(require_admin)):
    users = await User.find_all().sort("-created_at").to_list()

    user_ids = [str(u.id) for u in users]

    pipeline = [
        {"$match": {"user_id": {"$in": user_ids}}},
        {"$group": {"_id": "$user_id", "count": {"$sum": 1}}},
    ]
    upload_counts: dict[str, int] = {}
    async for doc in Defect.get_motor_collection().aggregate(pipeline):
        upload_counts[doc["_id"]] = doc["count"]

    log_pipeline = [
        {"$match": {"user_id": {"$in": user_ids}}},
        {"$group": {"_id": "$user_id", "last": {"$max": "$timestamp"}}},
    ]
    last_seen: dict[str, datetime] = {}
    async for doc in UserLog.get_motor_collection().aggregate(log_pipeline):
        uid = str(doc["_id"])
        ts = _as_utc_aware(doc.get("last"))
        if ts:
            last_seen[uid] = ts

    defect_last_pipeline = [
        {"$match": {"user_id": {"$in": user_ids}}},
        {"$group": {"_id": "$user_id", "last_upload": {"$max": "$created_at"}}},
    ]
    async for doc in Defect.get_motor_collection().aggregate(defect_last_pipeline):
        uid = str(doc["_id"])
        ts = _as_utc_aware(doc.get("last_upload"))
        if ts:
            prev = last_seen.get(uid)
            if prev is None or ts > prev:
                last_seen[uid] = ts

    logger.info("admin: fetched %d users", len(users))
    return [
        {
            "id": str(u.id),
            "email": u.email,
            "name": u.name,
            "role": u.role,
            "is_disabled": bool(u.is_disabled),
            "upload_count": upload_counts.get(str(u.id), 0),
            "last_activity": _iso_utc(last_seen.get(str(u.id)))
            or _iso_utc(_as_utc_aware(u.created_at)),
            "created_at": _iso_utc(_as_utc_aware(u.created_at)),
            "profile_photo": u.profile_photo,
        }
        for u in users
    ]


@router.post("/users/{user_id}/disable", status_code=status.HTTP_200_OK)
async def admin_disable_user_account(user_id: str, admin: User = Depends(require_admin)):
    if str(admin.id) == user_id:
        raise HTTPException(status_code=400, detail="You cannot disable your own account.")

    target = await _get_user_or_404(user_id)
    if target.is_disabled:
        raise HTTPException(status_code=400, detail="This account is already disabled.")

    if target.role == "admin" and _is_user_enabled(target):
        if await _count_enabled_admins() <= 1:
            raise HTTPException(
                status_code=400,
                detail="Cannot disable the last active administrator.",
            )

    target.is_disabled = True
    await target.save()
    await _log_admin_action(admin, "admin_disable_user", target=target)
    logger.info("admin: %s disabled user %s (%s)", admin.email, target.id, target.email)
    return {"ok": True, "is_disabled": True, "id": str(target.id)}


@router.post("/users/{user_id}/enable", status_code=status.HTTP_200_OK)
async def admin_enable_user_account(user_id: str, admin: User = Depends(require_admin)):
    target = await _get_user_or_404(user_id)
    if not target.is_disabled:
        raise HTTPException(status_code=400, detail="This account is already active.")

    target.is_disabled = False
    await target.save()
    await _log_admin_action(admin, "admin_enable_user", target=target)
    logger.info("admin: %s re-enabled user %s (%s)", admin.email, target.id, target.email)
    return {"ok": True, "is_disabled": False, "id": str(target.id)}


@router.delete("/users/{user_id}", status_code=status.HTTP_200_OK)
async def admin_delete_user_permanently(user_id: str, admin: User = Depends(require_admin)):
    if str(admin.id) == user_id:
        raise HTTPException(status_code=400, detail="You cannot delete your own account.")

    target = await _get_user_or_404(user_id)

    if target.role == "admin":
        if await _count_total_admins() <= 1:
            raise HTTPException(
                status_code=400,
                detail="Cannot delete the only administrator account.",
            )

    await _log_admin_action(admin, "admin_delete_user", target=target)

    n_defects = await _purge_user_defect_files(user_id)
    await _delete_user_profile_photo(target)

    # Remove activity logs where this user was the actor (not audit rows that reference them as target).
    await UserLog.get_motor_collection().delete_many({"user_id": user_id})

    await target.delete()

    logger.info(
        "admin: %s permanently deleted user %s (%s), defects_removed=%d",
        admin.email, user_id, target.email, n_defects,
    )
    return {"ok": True, "deleted_id": user_id, "defects_removed": n_defects}


@router.get("/users/{user_id}")
async def get_user_detail(user_id: str, _: User = Depends(require_admin)):
    user = await _get_user_or_404(user_id)
    return {
        "id": str(user.id),
        "email": user.email,
        "name": user.name,
        "role": user.role,
        "is_disabled": bool(user.is_disabled),
        "created_at": user.created_at.isoformat(),
        "mobile": user.mobile,
        "gender": user.gender,
        "age": user.age,
        "site": user.site,
        "location": user.location,
        "profile_photo": user.profile_photo,
    }


@router.post("/users/{user_id}/reset-password", status_code=status.HTTP_200_OK)
async def admin_reset_user_password(
    user_id: str,
    body: AdminResetPasswordBody,
    admin: User = Depends(require_admin),
):
    if str(admin.id) == user_id:
        raise HTTPException(
            status_code=400,
            detail="Change your own password from account settings, not from the admin console.",
        )
    target = await _get_user_or_404(user_id)
    target.password = hash_password(body.new_password)
    await target.save()
    await _log_admin_action(admin, "admin_reset_password", target=target)
    logger.info("admin: %s reset password for user %s", admin.email, target.email)
    return {"ok": True}


@router.patch("/users/{user_id}/role", status_code=status.HTTP_200_OK)
async def admin_set_user_role(user_id: str, body: AdminRoleBody, admin: User = Depends(require_admin)):
    if str(admin.id) == user_id:
        raise HTTPException(status_code=400, detail="You cannot change your own role.")

    target = await _get_user_or_404(user_id)
    new_role = body.role
    if new_role == target.role:
        return {"ok": True, "role": target.role}

    if target.role == "admin" and new_role == "user":
        if await _count_total_admins() <= 1:
            raise HTTPException(
                status_code=400,
                detail="Cannot demote the only administrator account.",
            )

    target.role = new_role
    await target.save()
    await _log_admin_action(admin, "admin_change_role", target=target)
    logger.info(
        "admin: %s set role for %s → %s",
        admin.email, target.email, new_role,
    )
    return {"ok": True, "role": target.role}


@router.get("/users/{user_id}/uploads")
async def get_user_uploads(user_id: str, _: User = Depends(require_admin)):
    defects = await Defect.find(Defect.user_id == user_id).sort("-created_at").to_list()
    logger.info("admin: fetched %d uploads for user %s", len(defects), user_id)
    return [_defect_to_dict(d) for d in defects]


@router.get("/uploads")
async def get_all_uploads(_: User = Depends(require_admin)):
    """Return every upload across all users, enriched with uploader identity."""
    defects = await Defect.find_all().sort("-created_at").to_list()
    user_ids = list({d.user_id for d in defects})
    users = await User.find(
        {"_id": {"$in": [PydanticObjectId(uid) for uid in user_ids]}}
    ).to_list()
    email_map = {str(u.id): u.email for u in users}
    name_map = {str(u.id): u.name for u in users}
    logger.info("admin: fetched all uploads (%d total)", len(defects))
    return [
        _defect_to_dict(
            d,
            email=email_map.get(d.user_id, "unknown"),
            name=name_map.get(d.user_id),
        )
        for d in defects
    ]


@router.post("/reports/analyze-item")
async def analyze_report_item(
    body: ReportAnalyzeItemBody,
    _: User = Depends(require_admin),
):
    try:
        defect = await Defect.get(PydanticObjectId(body.defect_id))
    except Exception:
        raise HTTPException(status_code=400, detail="Invalid defect ID") from None
    if not defect:
        raise HTTPException(status_code=404, detail="Upload not found")

    abs_img = _safe_image_abs_path(defect.image_path)
    if not abs_img:
        raise HTTPException(status_code=404, detail="Image file not found on storage")

    image_bytes = abs_img.read_bytes()
    ext = abs_img.suffix.lower()
    if ext == ".png":
        mime = "image/png"
    elif ext == ".webp":
        mime = "image/webp"
    else:
        mime = "image/jpeg"

    # Validation: never generate defect observations for unrelated images.
    if not await classify_construction_site_image(
        image_bytes=image_bytes,
        mime_type=mime,
    ):
        observation, recommendation = invalid_image_fields()
        return {
            "defect_id": str(defect.id),
            "category": defect.category or "Others",
            "reporting_date": defect.created_at.date().isoformat(),
            "tower": defect.tower,
            "floor": defect.floor,
            "flat": defect.flat,
            "room": defect.room,
            "image_path": defect.image_path,
            "observation": observation,
            "recommendation": recommendation,
        }

    prompt = build_executive_defect_report_prompt(
        description=defect.description or "",
        location=_defect_location_string(defect),
        issue_type=defect.category or "",
    )

    observation, recommendation = await generate_executive_defect_report(
        image_bytes=image_bytes,
        mime_type=mime,
        prompt=prompt,
    )

    return {
        "defect_id": str(defect.id),
        "category": defect.category or "Others",
        "reporting_date": defect.created_at.date().isoformat(),
        "tower": defect.tower,
        "floor": defect.floor,
        "flat": defect.flat,
        "room": defect.room,
        "image_path": defect.image_path,
        "observation": observation,
        "recommendation": recommendation,
    }


@router.post("/reports/generate-xlsx")
async def generate_report_xlsx(
    body: ReportGenerateXlsxBody,
    admin: User = Depends(require_admin),
):
    rows = await _collect_report_rows(body.entries)
    workbook = _build_report_workbook(rows=rows, base_url=body.base_url)
    filename = f"SiteSureLabs_Report_{datetime.now(timezone.utc).date().isoformat()}.xlsx"
    headers = {"Content-Disposition": f'attachment; filename="{filename}"'}

    await UserLog(
        user_id=str(admin.id),
        action=f"admin_generate_report:{len(rows)}",
    ).insert()

    return StreamingResponse(
        workbook,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


@router.post("/reports/generate-pptx")
async def generate_report_pptx(
    body: ReportGenerateXlsxBody,
    admin: User = Depends(require_admin),
):
    rows = await _collect_report_rows(body.entries)
    presentation = _build_report_presentation(rows=rows)
    filename = f"SiteSureLabs_Report_{datetime.now(timezone.utc).date().isoformat()}.pptx"
    headers = {"Content-Disposition": f'attachment; filename="{filename}"'}

    await UserLog(
        user_id=str(admin.id),
        action=f"admin_generate_report_pptx:{len(rows)}",
    ).insert()

    return StreamingResponse(
        presentation,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers=headers,
    )


@router.get("/logs")
async def get_logs(
    date_from: Optional[str] = Query(None, description="ISO date, e.g. 2025-01-01"),
    date_to: Optional[str] = Query(None, description="ISO date, e.g. 2025-12-31"),
    _: User = Depends(require_admin),
):
    query = {}

    if date_from or date_to:
        ts_filter = {}
        if date_from:
            try:
                ts_filter["$gte"] = datetime.fromisoformat(date_from).replace(tzinfo=timezone.utc)
            except ValueError:
                raise HTTPException(status_code=400, detail="Invalid date_from format")
        if date_to:
            try:
                end = datetime.fromisoformat(date_to).replace(
                    hour=23, minute=59, second=59, tzinfo=timezone.utc,
                )
                ts_filter["$lte"] = end
            except ValueError:
                raise HTTPException(status_code=400, detail="Invalid date_to format")
        query["timestamp"] = ts_filter

    logs = await UserLog.find(query).sort("-timestamp").to_list()

    user_ids = list({log.user_id for log in logs})
    users = await User.find({"_id": {"$in": [PydanticObjectId(uid) for uid in user_ids]}}).to_list()
    email_map = {str(u.id): u.email for u in users}
    name_map = {str(u.id): u.name for u in users}

    logger.info("admin: fetched %d log entries", len(logs))
    return [
        {
            "id": str(log.id),
            "user_id": log.user_id,
            "email": email_map.get(log.user_id, "unknown"),
            "name": name_map.get(log.user_id),
            "action": log.action,
            "timestamp": log.timestamp.isoformat(),
            "target_user_id": log.target_user_id,
            "target_email": log.target_email,
        }
        for log in logs
    ]


@router.delete("/logs")
async def clear_all_logs(admin: User = Depends(require_admin)):
    """Remove every activity log entry (admin-only)."""
    result = await UserLog.get_motor_collection().delete_many({})
    deleted = int(result.deleted_count or 0)
    logger.info("admin: cleared all activity logs (%d) by=%s", deleted, admin.email)
    return {"ok": True, "deleted": deleted}


_BULK_DELETE_IDS_MAX = 8000


@router.post("/uploads/bulk-delete", status_code=status.HTTP_200_OK)
async def bulk_delete_uploads(body: BulkDeleteUploadsBody, admin: User = Depends(require_admin)):
    """Remove many defect uploads (disk + DB). Uses best-effort file deletes so rows still clear if a file is stuck."""
    removed = 0

    if body.mode == "all":
        defects = await Defect.find_all().to_list()
        for defect in defects:
            await _delete_defect_upload_best_effort(defect)
            removed += 1
    elif body.mode == "user":
        uid = (body.user_id or "").strip()
        if not uid:
            raise HTTPException(status_code=400, detail="user_id is required for mode=user")
        removed = await _purge_user_defect_files(uid)
    else:
        raw_ids = body.defect_ids
        if not raw_ids:
            raise HTTPException(status_code=400, detail="defect_ids is required for mode=ids")
        if len(raw_ids) > _BULK_DELETE_IDS_MAX:
            raise HTTPException(
                status_code=400,
                detail=f"Too many defect IDs (maximum {_BULK_DELETE_IDS_MAX})",
            )
        try:
            oid_list = [PydanticObjectId(x) for x in raw_ids]
        except Exception:
            raise HTTPException(status_code=400, detail="Invalid defect ID in defect_ids") from None
        defects = await Defect.find({"_id": {"$in": oid_list}}).to_list()
        for defect in defects:
            await _delete_defect_upload_best_effort(defect)
            removed += 1

    await UserLog(
        user_id=str(admin.id),
        action=f"bulk_delete_uploads:{body.mode}:{removed}",
    ).insert()
    logger.info(
        "admin: bulk_delete_uploads mode=%s removed=%d by=%s",
        body.mode,
        removed,
        admin.email,
    )
    return {"ok": True, "removed": removed, "mode": body.mode}


@router.delete("/uploads/{upload_id}")
async def delete_upload(upload_id: str, admin: User = Depends(require_admin)):
    try:
        defect = await Defect.get(PydanticObjectId(upload_id))
    except Exception:
        raise HTTPException(status_code=400, detail="Invalid upload ID")
    if not defect:
        raise HTTPException(status_code=404, detail="Upload not found")

    had_file = (REPO_ROOT / defect.image_path).exists()
    await _delete_defect_upload_strict(defect)

    await UserLog(
        user_id=str(admin.id),
        action=f"delete_upload:{upload_id}",
    ).insert()

    logger.info(
        "admin: deleted upload %s (had_file=%s, by=%s)",
        upload_id, had_file, admin.email,
    )
    return {"deleted": True, "id": upload_id}


def _start_of_utc_today() -> datetime:
    now = datetime.now(timezone.utc)
    return now.replace(hour=0, minute=0, second=0, microsecond=0)


def _parse_analytics_date_range(date: str, tz_offset: int) -> tuple[datetime, datetime]:
    """Map a calendar date + JS getTimezoneOffset() to a UTC query window."""
    try:
        day = datetime.strptime(date, "%Y-%m-%d").date()
    except ValueError as exc:
        raise HTTPException(
            status_code=400,
            detail="Invalid date; use YYYY-MM-DD",
        ) from exc
    start = datetime(day.year, day.month, day.day, tzinfo=timezone.utc) + timedelta(
        minutes=tz_offset
    )
    return start, start + timedelta(days=1)


def _analytics_area_label(d: Defect) -> str:
    parts: list[str] = []
    if d.tower:
        parts.append(f"Tower {d.tower}")
    if d.floor:
        parts.append(f"Floor {d.floor}")
    if d.flat:
        parts.append(f"Flat {d.flat}")
    if d.room:
        parts.append(d.room)
    return " · ".join(parts) if parts else "Unknown area"


def _tower_floor_label(tower: str, floor: str) -> str:
    """Human-readable tower + floor, e.g. 'Tower C, Floor 15' (not 'Tower C · 15')."""
    t = (tower or "").strip() or "Unknown area"
    f = (floor or "").strip()
    if not f or f in ("—", "-"):
        return t
    if f.lower().startswith("floor"):
        return f"{t}, {f}"
    return f"{t}, Floor {f}"


def _analytics_category_label(d: Defect) -> str:
    cat = (d.category or "").strip()
    if cat:
        return cat
    desc = (d.description or "").strip()
    if desc:
        return desc if len(desc) <= 72 else f"{desc[:69]}…"
    return "Uncategorized"


def _top_counts(items: list[tuple[str, int]], *, limit: int = 6) -> list[dict]:
    ranked = sorted(items, key=lambda pair: (-pair[1], pair[0].lower()))
    return [{"name": name, "count": count} for name, count in ranked[:limit]]


_MITIGATION_HINTS: list[tuple[tuple[str, ...], str]] = [
    (
        ("crack", "structural", "settlement", "column", "beam"),
        "Schedule structural punch-list review on affected towers and verify closure with engineer sign-off.",
    ),
    (
        ("water", "leak", "seep", "damp", "moisture", "plumb", "pipe"),
        "Prioritize waterproofing and plumbing rectification; re-inspect wet areas after rework.",
    ),
    (
        ("electrical", "wiring", "switch", "cable", "mcb"),
        "Coordinate with MEP for electrical safety checks and targeted fixes in hotspot flats.",
    ),
    (
        ("paint", "finish", "plaster", "tile", "grout", "floor"),
        "Deploy finishing teams to high-volume rooms and verify quality with photo evidence.",
    ),
    (
        ("door", "window", "frame", "glass", "hardware"),
        "Batch carpentry and fenestration rework on the most affected floors.",
    ),
]


def _pick_mitigation(by_category: list[dict]) -> str:
    haystack = " ".join(item["name"].lower() for item in by_category[:5])
    for keywords, hint in _MITIGATION_HINTS:
        if any(word in haystack for word in keywords):
            return hint
    return (
        "Cluster recurring issues by tower and room, assign owners, and track closure with photo verification."
    )


def _build_analytics_insight(
    *,
    observation_count: int,
    by_project: list[dict],
    by_tower: list[dict],
    by_area: list[dict],
    by_category: list[dict],
) -> str:
    if observation_count == 0:
        return (
            "No observations were recorded on this date. Try another day or encourage site teams "
            "to capture uploads from critical walkthroughs."
        )

    parts: list[str] = []
    if by_project:
        top = by_project[0]
        parts.append(
            f"{top['name']} led activity with {top['count']} observation"
            f"{'' if top['count'] == 1 else 's'}."
        )
    if by_tower:
        top = by_tower[0]
        parts.append(
            f"Tower {top['name']} showed the highest defect concentration ({top['count']})."
        )
    if by_area:
        top = by_area[0]
        parts.append(f"The busiest hotspot was {top['name']} ({top['count']} issues).")
    if by_category:
        top = by_category[0]
        parts.append(
            f"The most frequent issue type was {top['name']} ({top['count']} recorded)."
        )
    parts.append(f"Recommended action: {_pick_mitigation(by_category)}")
    return " ".join(parts)


@router.get("/analytics")
async def get_analytics(
    date: str = Query(..., description="Calendar date YYYY-MM-DD"),
    tz_offset: int = Query(
        0,
        description="Minutes from JS Date.getTimezoneOffset() for local-day boundaries",
    ),
    _: User = Depends(require_admin),
):
    start, end = _parse_analytics_date_range(date, tz_offset)
    defects = await Defect.find(
        Defect.created_at >= start,
        Defect.created_at < end,
    ).to_list()

    user_ids = list({d.user_id for d in defects})
    if user_ids:
        users = await User.find(
            {"_id": {"$in": [PydanticObjectId(uid) for uid in user_ids]}}
        ).to_list()
    else:
        users = []
    email_map = {str(u.id): u.email for u in users}
    name_map = {str(u.id): u.name for u in users}

    uploader_counts: dict[str, int] = {}
    project_counts: dict[str, int] = {}
    tower_counts: dict[str, int] = {}
    area_counts: dict[str, int] = {}
    category_counts: dict[str, int] = {}

    for d in defects:
        uploader_counts[d.user_id] = uploader_counts.get(d.user_id, 0) + 1

        project = (getattr(d, "project", "") or "").strip() or "Unassigned project"
        project_counts[project] = project_counts.get(project, 0) + 1

        tower = (d.tower or "").strip() or "Unknown"
        tower_counts[tower] = tower_counts.get(tower, 0) + 1

        area = _analytics_area_label(d)
        area_counts[area] = area_counts.get(area, 0) + 1

        category = _analytics_category_label(d)
        category_counts[category] = category_counts.get(category, 0) + 1

    uploaders = sorted(
        [
            {
                "user_id": uid,
                "email": email_map.get(uid, "unknown"),
                "name": name_map.get(uid),
                "count": count,
            }
            for uid, count in uploader_counts.items()
        ],
        key=lambda row: (-row["count"], (row["email"] or "").lower()),
    )

    by_project = _top_counts(list(project_counts.items()))
    by_tower = _top_counts(list(tower_counts.items()))
    by_area = _top_counts(list(area_counts.items()))
    by_category = _top_counts(list(category_counts.items()))

    insight = _build_analytics_insight(
        observation_count=len(defects),
        by_project=by_project,
        by_tower=by_tower,
        by_area=by_area,
        by_category=by_category,
    )

    logger.info(
        "admin: analytics for %s → %d observations, %d uploaders",
        date,
        len(defects),
        len(uploaders),
    )

    return {
        "date": date,
        "observation_count": len(defects),
        "uploader_count": len(uploaders),
        "uploaders": uploaders,
        "by_project": by_project,
        "by_tower": by_tower,
        "by_area": by_area,
        "by_category": by_category,
        "insight": insight,
    }


_CRITICAL_KEYWORDS = (
    "structural",
    "crack",
    "collapse",
    "safety",
    "fire",
    "electrical",
    "shock",
    "gas",
    "seepage",
    "foundation",
    "hazard",
    "urgent",
    "critical",
    "leak",
    "waterproof",
)


def _utc_day_key(dt: datetime, tz_offset_min: int) -> str:
    aware = _as_utc_aware(dt)
    if aware is None:
        return ""
    local = aware + timedelta(minutes=tz_offset_min)
    return local.strftime("%Y-%m-%d")


def _is_critical_defect(d: Defect) -> bool:
    hay = f"{(d.category or '')} {(d.description or '')}".lower()
    return any(word in hay for word in _CRITICAL_KEYWORDS)


def _format_activity_action(action: str) -> str:
    if action == "upload":
        return "Image uploaded"
    if action == "register":
        return "New user registered"
    if action == "login":
        return "User signed in"
    if action == "admin_login":
        return "Admin signed in"
    if action.startswith("admin_generate_report_pptx:"):
        n = action.split(":", 1)[-1]
        return f"PowerPoint report generated ({n} items)"
    if action.startswith("admin_generate_report:"):
        n = action.split(":", 1)[-1]
        return f"Excel report generated ({n} items)"
    if action.startswith("bulk_delete_uploads:"):
        return "Bulk upload cleanup"
    if action.startswith("delete_upload:"):
        return "Upload removed"
    if action.startswith("admin_"):
        return action.replace("admin_", "Admin ").replace("_", " ")
    return action.replace("_", " ")


@router.get("/workspace")
async def get_workspace_dashboard(
    tz_offset: int = Query(
        0,
        description="Minutes from JS Date.getTimezoneOffset() for local-day boundaries",
    ),
    _: User = Depends(require_admin),
):
    """Aggregated real-time analytics for the admin Stats workspace tab."""
    now = datetime.now(timezone.utc)
    start_today = _start_of_utc_today()
    week_ago = now - timedelta(days=7)
    month_ago = now - timedelta(days=30)

    total_users = await User.count()
    disabled_users = await User.find(User.is_disabled == True).count()
    total_uploads = await Defect.count()
    total_logs = await UserLog.count()
    users_today = await User.find(User.created_at >= start_today).count()
    users_week = await User.find(User.created_at >= week_ago).count()
    uploads_today = await Defect.find(Defect.created_at >= start_today).count()
    uploads_week = await Defect.find(Defect.created_at >= week_ago).count()
    logs_today = await UserLog.find(UserLog.timestamp >= start_today).count()
    logs_week = await UserLog.find(UserLog.timestamp >= week_ago).count()

    defects_30d = await Defect.find(Defect.created_at >= month_ago).to_list()
    logs_30d = await UserLog.find(UserLog.timestamp >= month_ago).to_list()
    logs_recent = await UserLog.find().sort("-timestamp").limit(40).to_list()

    user_ids = list(
        {d.user_id for d in defects_30d}
        | {log.user_id for log in logs_recent}
    )
    if user_ids:
        users = await User.find(
            {"_id": {"$in": [PydanticObjectId(uid) for uid in user_ids]}}
        ).to_list()
    else:
        users = []
    email_map = {str(u.id): u.email for u in users}
    name_map = {str(u.id): u.name for u in users}

    daily_keys = [
        _utc_day_key(now - timedelta(days=i), tz_offset)
        for i in range(13, -1, -1)
    ]
    daily_upload_map = {k: 0 for k in daily_keys}
    for d in defects_30d:
        key = _utc_day_key(_as_utc_aware(d.created_at) or now, tz_offset)
        if key in daily_upload_map:
            daily_upload_map[key] += 1
    upload_trend_daily = [
        {"date": k, "label": datetime.strptime(k, "%Y-%m-%d").strftime("%b %d"), "count": daily_upload_map[k]}
        for k in daily_keys
    ]

    upload_trend_weekly: list[dict] = []
    for i in range(7, -1, -1):
        window_end = now - timedelta(days=i * 7)
        window_start = window_end - timedelta(days=7)
        count = sum(
            1
            for d in defects_30d
            if window_start <= _as_utc_aware(d.created_at) < window_end
        )
        upload_trend_weekly.append(
            {
                "week_start": _utc_day_key(window_start, tz_offset),
                "label": "This wk" if i == 0 else f"-{i}w",
                "count": count,
            }
        )

    category_counts: dict[str, int] = {}
    tower_counts: dict[str, int] = {}
    floor_counts: dict[str, int] = {}
    tower_floor_counts: dict[tuple[str, str], int] = {}
    heat_cells: dict[tuple[str, str], int] = {}
    critical_items: list[dict] = []
    uploader_7d: dict[str, dict] = {}

    for d in defects_30d:
        created_at = _as_utc_aware(d.created_at)
        cat = _analytics_category_label(d)
        category_counts[cat] = category_counts.get(cat, 0) + 1

        tower = (d.tower or "").strip() or "—"
        floor = (d.floor or "").strip() or "—"
        tower_counts[tower] = tower_counts.get(tower, 0) + 1
        floor_counts[floor] = floor_counts.get(floor, 0) + 1
        tf_key = (tower, floor)
        tower_floor_counts[tf_key] = tower_floor_counts.get(tf_key, 0) + 1
        heat_cells[tf_key] = heat_cells.get(tf_key, 0) + 1

        if _is_critical_defect(d):
            critical_items.append(
                {
                    "id": str(d.id),
                    "label": _analytics_category_label(d),
                    "area": _analytics_area_label(d),
                    "at": _iso_utc(created_at) or "",
                }
            )

        if created_at is not None and created_at >= week_ago:
            uid = d.user_id
            created_iso = _iso_utc(created_at) or ""
            row = uploader_7d.get(uid) or {
                "user_id": uid,
                "email": email_map.get(uid, "unknown"),
                "name": name_map.get(uid),
                "count": 0,
                "last_at": created_iso,
            }
            row["count"] += 1
            if created_iso > row["last_at"]:
                row["last_at"] = created_iso
            uploader_7d[uid] = row

    by_category = _top_counts(list(category_counts.items()), limit=8)
    tower_floor = _top_counts(
        [(_tower_floor_label(t[0], t[1]), c) for t, c in tower_floor_counts.items()],
        limit=8,
    )

    top_towers = [t for t, _ in sorted(tower_counts.items(), key=lambda x: (-x[1], x[0]))[:6]]
    top_floors = [f for f, _ in sorted(floor_counts.items(), key=lambda x: (-x[1], x[0]))[:6]]
    if not top_towers:
        top_towers = ["—"]
    if not top_floors:
        top_floors = ["—"]
    heatmap = {
        "towers": top_towers,
        "floors": top_floors,
        "cells": [
            {
                "tower": tower,
                "floor": floor,
                "count": heat_cells.get((tower, floor), 0),
            }
            for tower in top_towers
            for floor in top_floors
        ],
        "max": max(heat_cells.values()) if heat_cells else 1,
    }

    report_daily_map = {k: 0 for k in daily_keys}
    reports_week = 0
    reports_today = 0
    for log in logs_30d:
        if not log.action.startswith("admin_generate_report"):
            continue
        log_ts = _as_utc_aware(log.timestamp)
        if log_ts is None:
            continue
        key = _utc_day_key(log_ts, tz_offset)
        if key in report_daily_map:
            report_daily_map[key] += 1
        if log_ts >= week_ago:
            reports_week += 1
        if log_ts >= start_today:
            reports_today += 1
    report_trend_daily = [
        {"date": k, "label": datetime.strptime(k, "%Y-%m-%d").strftime("%b %d"), "count": report_daily_map[k]}
        for k in daily_keys
    ]

    user_activity = sorted(
        uploader_7d.values(),
        key=lambda row: (-row["count"], (row["email"] or "").lower()),
    )[:8]

    critical_items.sort(key=lambda row: row["at"], reverse=True)
    critical_count = len(critical_items)
    critical_recent = critical_items[:6]

    timeline: list[dict] = []
    for log in logs_recent[:25]:
        timeline.append(
            {
                "type": "system",
                "label": _format_activity_action(log.action),
                "actor": email_map.get(log.user_id, "unknown"),
                "at": _iso_utc(_as_utc_aware(log.timestamp)) or "",
            }
        )
    for d in sorted(
        defects_30d,
        key=lambda x: _as_utc_aware(x.created_at) or datetime.min.replace(tzinfo=timezone.utc),
        reverse=True,
    )[:12]:
        uid = d.user_id
        timeline.append(
            {
                "type": "upload",
                "label": _analytics_category_label(d),
                "actor": email_map.get(uid, "unknown"),
                "at": _iso_utc(_as_utc_aware(d.created_at)) or "",
            }
        )
    timeline.sort(key=lambda row: row["at"], reverse=True)
    timeline = timeline[:20]

    ai_success = reports_week
    ai_estimate_fail = max(0, uploads_week - reports_week)

    return {
        "stats": {
            "total_users": total_users,
            "disabled_users": disabled_users,
            "total_uploads": total_uploads,
            "total_logs": total_logs,
            "users_today": users_today,
            "users_week": users_week,
            "uploads_today": uploads_today,
            "uploads_week": uploads_week,
            "logs_today": logs_today,
            "logs_week": logs_week,
        },
        "kpis": {
            "critical_count": critical_count,
            "uploads_today": uploads_today,
            "reports_today": reports_today,
            "active_uploaders_7d": len(uploader_7d),
            "ai_success_week": ai_success,
            "ai_pending_estimate": ai_estimate_fail,
        },
        "upload_trend_daily": upload_trend_daily,
        "upload_trend_weekly": upload_trend_weekly,
        "by_category": by_category,
        "heatmap": heatmap,
        "tower_floor": tower_floor,
        "report_trend_daily": report_trend_daily,
        "user_activity": user_activity,
        "critical": {"count": critical_count, "recent": critical_recent},
        "ai_metrics": {
            "reports_today": reports_today,
            "reports_week": reports_week,
            "uploads_week": uploads_week,
            "success_week": ai_success,
            "pending_estimate": ai_estimate_fail,
        },
        "timeline": timeline,
        "generated_at": now.isoformat(),
    }


@router.get("/stats")
async def get_stats(_: User = Depends(require_admin)):
    total_users = await User.count()
    disabled_users = await User.find(User.is_disabled == True).count()
    total_uploads = await Defect.count()
    total_logs = await UserLog.count()

    start_today = _start_of_utc_today()
    week_ago = datetime.now(timezone.utc) - timedelta(days=7)

    users_today = await User.find(User.created_at >= start_today).count()
    users_week = await User.find(User.created_at >= week_ago).count()
    uploads_today = await Defect.find(Defect.created_at >= start_today).count()
    uploads_week = await Defect.find(Defect.created_at >= week_ago).count()
    logs_today = await UserLog.find(UserLog.timestamp >= start_today).count()
    logs_week = await UserLog.find(UserLog.timestamp >= week_ago).count()

    return {
        "total_users": total_users,
        "disabled_users": disabled_users,
        "total_uploads": total_uploads,
        "total_logs": total_logs,
        "users_today": users_today,
        "users_week": users_week,
        "uploads_today": uploads_today,
        "uploads_week": uploads_week,
        "logs_today": logs_today,
        "logs_week": logs_week,
    }


@router.get("/debug/uploads")
async def debug_uploads(_: User = Depends(require_admin)):
    """Temporary debug endpoint: returns every Defect with file-existence check."""
    defects = await Defect.find_all().sort("-created_at").to_list()
    result = []
    for d in defects:
        file_abs = UPLOADS_DIR.parent / d.image_path
        result.append({
            "id": str(d.id),
            "user_id": d.user_id,
            "image_path": d.image_path,
            "file_exists": file_abs.exists(),
            "file_size": file_abs.stat().st_size if file_abs.exists() else None,
            "created_at": d.created_at.isoformat(),
        })
    return result


@router.get("/debug/storage")
async def debug_storage(_: User = Depends(require_admin)):
    """Temporary debug endpoint: lists physical files on disk under uploads/."""
    files = []
    for p in sorted(UPLOADS_DIR.rglob("*")):
        if p.is_file():
            files.append({
                "path": str(p.relative_to(UPLOADS_DIR.parent)),
                "size": p.stat().st_size,
            })
    return {"upload_dir": str(UPLOADS_DIR), "file_count": len(files), "files": files}
